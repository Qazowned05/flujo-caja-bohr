import uuid
from calendar import monthrange
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Annotated, Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import get_db
from app.modules.ganancias_perdidas.models import (
    AsientoResultado,
    CentroResultado,
    LoteResultado,
    NaturalezaRubro,
    OrigenAsientoResultado,
    ReglaDistribucion,
    ReglaDistribucionLinea,
    RubroResultado,
)
from app.modules.ganancias_perdidas.schemas import (
    AnularAsientoResultado,
    AsientoResultadoCreate,
    AsientoResultadoRead,
    AsientoResultadoUpdate,
    CentroResultadoCreate,
    CentroResultadoRead,
    CentroResultadoUpdate,
    ImportResultadoRead,
    ReglaDistribucionCreate,
    ReglaDistribucionRead,
    RentabilidadCentroResultadoRead,
    ResultadoRubroRead,
    ResumenGananciasPerdidasRead,
    RubroResultadoCreate,
    RubroResultadoRead,
    RubroResultadoUpdate,
)
from app.modules.shared.audit import log_audit
from app.modules.shared.models import AuditAction
from app.modules.usuarios.deps import get_current_user, require_admin
from app.modules.usuarios.models import User

router = APIRouter(prefix="/ganancias-perdidas", tags=["ganancias-perdidas"])
ZERO = Decimal("0")
TEMPLATE_HEADERS = [
    "periodo",
    "rubro_codigo",
    "monto_sin_igv",
    "centro_codigo",
    "descripcion",
    "observaciones",
]


def require_enabled() -> None:
    if not settings.profit_and_loss_enabled:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Modulo no habilitado.")


def require_record(db: Session, model: Any, record_id: uuid.UUID) -> Any:
    record = db.get(model, record_id)
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Registro no encontrado.")
    return record


def require_active(db: Session, model: Any, record_id: uuid.UUID, label: str) -> Any:
    record = require_record(db, model, record_id)
    if not record.is_active:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"{label} activo no encontrado.",
        )
    return record


def save_change(
    db: Session, user: User, record: Any, action: AuditAction, changes: dict[str, Any]
) -> None:
    log_audit(
        db,
        user=user,
        table=record.__tablename__,
        record_id=record.id,
        action=action,
        changes=changes,
    )
    db.commit()
    db.refresh(record)


def apply_update(record: Any, values: dict[str, Any]) -> dict[str, Any]:
    changes = {
        field: {"old": getattr(record, field), "new": value}
        for field, value in values.items()
        if getattr(record, field) != value
    }
    for field, value in values.items():
        setattr(record, field, value)
    return changes


def active_statement(model: Any, include_inactive: bool):
    statement = select(model)
    return statement if include_inactive else statement.where(model.is_active.is_(True))


@router.get(
    "/centros", response_model=list[CentroResultadoRead], dependencies=[Depends(require_enabled)]
)
def list_centros(
    include_inactive: bool = False,
    _: Annotated[User, Depends(require_admin)] = None,
    db: Annotated[Session, Depends(get_db)] = None,
):
    return list(
        db.scalars(
            active_statement(CentroResultado, include_inactive).order_by(
                CentroResultado.orden, CentroResultado.nombre
            )
        )
    )


@router.post(
    "/centros",
    response_model=CentroResultadoRead,
    status_code=status.HTTP_201_CREATED,
    dependencies=[Depends(require_enabled)],
)
def create_centro(
    data: CentroResultadoCreate,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    values = data.model_dump()
    if db.scalar(
        select(CentroResultado).where(
            (CentroResultado.codigo == values["codigo"])
            | (CentroResultado.nombre == values["nombre"])
        )
    ):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="El codigo o nombre de linea de negocio ya existe.",
        )
    record = CentroResultado(**values)
    db.add(record)
    db.flush()
    save_change(db, user, record, AuditAction.CREATE, values)
    return record


@router.patch(
    "/centros/{centro_id}",
    response_model=CentroResultadoRead,
    dependencies=[Depends(require_enabled)],
)
def update_centro(
    centro_id: uuid.UUID,
    data: CentroResultadoUpdate,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    record = require_record(db, CentroResultado, centro_id)
    values = data.model_dump(exclude_unset=True)
    for field in ("codigo", "nombre"):
        if field in values and db.scalar(
            select(CentroResultado).where(
                getattr(CentroResultado, field) == values[field], CentroResultado.id != centro_id
            )
        ):
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT, detail=f"El {field} ya existe."
            )
    changes = apply_update(record, values)
    if changes:
        save_change(db, user, record, AuditAction.UPDATE, changes)
    return record


@router.delete(
    "/centros/{centro_id}",
    response_model=CentroResultadoRead,
    dependencies=[Depends(require_enabled)],
)
def delete_centro(
    centro_id: uuid.UUID,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    record = require_active(db, CentroResultado, centro_id, "Linea de negocio")
    record.is_active, record.deleted_at = False, datetime.now(UTC)
    save_change(
        db, user, record, AuditAction.SOFT_DELETE, {"is_active": {"old": True, "new": False}}
    )
    return record


@router.get(
    "/rubros", response_model=list[RubroResultadoRead], dependencies=[Depends(require_enabled)]
)
def list_rubros(
    include_inactive: bool = False,
    _: Annotated[User, Depends(require_admin)] = None,
    db: Annotated[Session, Depends(get_db)] = None,
):
    return list(
        db.scalars(
            active_statement(RubroResultado, include_inactive).order_by(
                RubroResultado.orden, RubroResultado.nombre
            )
        )
    )


@router.post(
    "/rubros",
    response_model=RubroResultadoRead,
    status_code=status.HTTP_201_CREATED,
    dependencies=[Depends(require_enabled)],
)
def create_rubro(
    data: RubroResultadoCreate,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    values = data.model_dump()
    if db.scalar(select(RubroResultado).where(RubroResultado.codigo == values["codigo"])):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT, detail="El codigo de rubro ya existe."
        )
    if values["padre_id"]:
        require_active(db, RubroResultado, values["padre_id"], "Rubro padre")
    record = RubroResultado(**values)
    db.add(record)
    db.flush()
    save_change(db, user, record, AuditAction.CREATE, values)
    return record


@router.patch(
    "/rubros/{rubro_id}", response_model=RubroResultadoRead, dependencies=[Depends(require_enabled)]
)
def update_rubro(
    rubro_id: uuid.UUID,
    data: RubroResultadoUpdate,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    record = require_record(db, RubroResultado, rubro_id)
    values = data.model_dump(exclude_unset=True)
    if "codigo" in values and db.scalar(
        select(RubroResultado).where(
            RubroResultado.codigo == values["codigo"], RubroResultado.id != rubro_id
        )
    ):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT, detail="El codigo de rubro ya existe."
        )
    if values.get("padre_id"):
        if values["padre_id"] == rubro_id:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Un rubro no puede ser su propio padre.",
            )
        require_active(db, RubroResultado, values["padre_id"], "Rubro padre")
    changes = apply_update(record, values)
    if changes:
        save_change(db, user, record, AuditAction.UPDATE, changes)
    return record


@router.delete(
    "/rubros/{rubro_id}", response_model=RubroResultadoRead, dependencies=[Depends(require_enabled)]
)
def delete_rubro(
    rubro_id: uuid.UUID,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    record = require_active(db, RubroResultado, rubro_id, "Rubro")
    if db.scalar(select(AsientoResultado.id).where(AsientoResultado.rubro_id == rubro_id).limit(1)):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="No se puede inactivar un rubro con asientos.",
        )
    record.is_active, record.deleted_at = False, datetime.now(UTC)
    save_change(
        db, user, record, AuditAction.SOFT_DELETE, {"is_active": {"old": True, "new": False}}
    )
    return record


def validate_distribution(db: Session, values: ReglaDistribucionCreate) -> None:
    if values.rubro_id:
        require_active(db, RubroResultado, values.rubro_id, "Rubro")
    for line in values.lineas:
        require_active(db, CentroResultado, line.centro_resultado_id, "Linea de negocio")


@router.get(
    "/reglas-distribucion",
    response_model=list[ReglaDistribucionRead],
    dependencies=[Depends(require_enabled)],
)
def list_distribution_rules(
    include_inactive: bool = False,
    _: Annotated[User, Depends(require_admin)] = None,
    db: Annotated[Session, Depends(get_db)] = None,
):
    return list(
        db.scalars(
            active_statement(ReglaDistribucion, include_inactive).order_by(ReglaDistribucion.nombre)
        )
    )


@router.post(
    "/reglas-distribucion",
    response_model=ReglaDistribucionRead,
    status_code=status.HTTP_201_CREATED,
    dependencies=[Depends(require_enabled)],
)
def create_distribution_rule(
    data: ReglaDistribucionCreate,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    if db.scalar(select(ReglaDistribucion).where(ReglaDistribucion.nombre == data.nombre)):
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="La regla ya existe.")
    validate_distribution(db, data)
    values = data.model_dump(exclude={"lineas"})
    record = ReglaDistribucion(**values)
    record.lineas = [ReglaDistribucionLinea(**line.model_dump()) for line in data.lineas]
    db.add(record)
    db.flush()
    save_change(db, user, record, AuditAction.CREATE, data.model_dump(mode="json"))
    return record


@router.patch(
    "/reglas-distribucion/{regla_id}",
    response_model=ReglaDistribucionRead,
    dependencies=[Depends(require_enabled)],
)
def update_distribution_rule(
    regla_id: uuid.UUID,
    data: ReglaDistribucionCreate,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    record = require_record(db, ReglaDistribucion, regla_id)
    duplicate = db.scalar(
        select(ReglaDistribucion).where(
            ReglaDistribucion.nombre == data.nombre, ReglaDistribucion.id != regla_id
        )
    )
    if duplicate:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="La regla ya existe.")
    validate_distribution(db, data)
    changes = apply_update(record, data.model_dump(exclude={"lineas"}))
    previous_lines = [
        {"centro_resultado_id": str(line.centro_resultado_id), "porcentaje": str(line.porcentaje)}
        for line in record.lineas
    ]
    # Flush orphaned lines before inserting replacements to preserve the unique rule/center pair.
    record.lineas.clear()
    db.flush()
    record.lineas = [ReglaDistribucionLinea(**line.model_dump()) for line in data.lineas]
    changes["lineas"] = {"old": previous_lines, "new": data.model_dump(mode="json")["lineas"]}
    save_change(db, user, record, AuditAction.UPDATE, changes)
    return record


@router.delete(
    "/reglas-distribucion/{regla_id}",
    response_model=ReglaDistribucionRead,
    dependencies=[Depends(require_enabled)],
)
def delete_distribution_rule(
    regla_id: uuid.UUID,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    record = require_active(db, ReglaDistribucion, regla_id, "Regla")
    record.is_active, record.deleted_at = False, datetime.now(UTC)
    save_change(
        db, user, record, AuditAction.SOFT_DELETE, {"is_active": {"old": True, "new": False}}
    )
    return record


@router.get("/plantilla.xlsx", dependencies=[Depends(require_enabled)])
def download_template(
    _: Annotated[User, Depends(require_admin)], db: Annotated[Session, Depends(get_db)]
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Asientos"
    sheet.append(TEMPLATE_HEADERS)
    sheet.append(
        [
            "2026-06",
            "VENTAS",
            50000,
            "200",
            "Ventas junio",
            "",
        ]
    )
    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)
    sheet.freeze_panes = "A2"
    for column, width in {
        "A": 14,
        "B": 28,
        "C": 18,
        "D": 18,
        "E": 35,
        "F": 35,
    }.items():
        sheet.column_dimensions[column].width = width

    centers = list(
        db.scalars(
            select(CentroResultado)
            .where(CentroResultado.is_active.is_(True))
            .order_by(CentroResultado.orden)
        )
    )
    rubros = list(
        db.scalars(
            select(RubroResultado)
            .where(RubroResultado.is_active.is_(True))
            .order_by(RubroResultado.orden)
        )
    )
    centers_sheet = workbook.create_sheet("Lineas de negocio")
    centers_sheet.append(["codigo", "nombre"])
    for center in centers:
        centers_sheet.append([center.codigo, center.nombre])
    parent_ids = {rubro.padre_id for rubro in rubros if rubro.padre_id}
    rubros = [rubro for rubro in rubros if rubro.id not in parent_ids]
    rubros_sheet = workbook.create_sheet("Rubros")
    rubros_sheet.append(["codigo", "nombre", "naturaleza", "rubro_padre"])
    rubro_by_id = {rubro.id: rubro.codigo for rubro in rubros}
    for rubro in rubros:
        rubros_sheet.append([rubro.codigo, rubro.nombre, rubro.naturaleza.value, rubro_by_id.get(rubro.padre_id, "")])
    rules_sheet = workbook.create_sheet("Reglas distribucion")
    rules_sheet.append(
        [
            "regla",
            "rubro",
            "centro",
            "porcentaje",
            "vigente_desde",
            "vigente_hasta",
        ]
    )
    for rule in db.scalars(select(ReglaDistribucion).where(ReglaDistribucion.is_active.is_(True))):
        for line in rule.lineas:
            rules_sheet.append(
                [
                    rule.nombre,
                    rubro_by_id.get(rule.rubro_id, ""),
                    next((center.codigo for center in centers if center.id == line.centro_resultado_id), ""),
                    line.porcentaje,
                    rule.vigente_desde,
                    rule.vigente_hasta,
                ]
            )
    instructions = workbook.create_sheet("Guia de tipificacion")
    instructions.append(["Campo", "Uso"])
    instructions.append(["periodo", "Escribe solo mes y ano: 2026-06. El sistema usa el ultimo dia del mes."])
    instructions.append(["monto_sin_igv", "Escribe siempre un monto positivo, sin IGV."])
    instructions.append(["centro_codigo", "Selecciona una linea para importes directos. Dejalo vacio para repartir un gasto compartido."])
    instructions.append(
        ["rubro_codigo", "Obligatorio. Selecciona el codigo en la lista desplegable."]
    )
    instructions.append(["Calculo", "El sistema identifica automaticamente si el importe es ingreso, descuento, costo o gasto."])
    instructions.append(["Guias", "Las hojas Lineas de negocio, Rubros y Reglas distribucion son solo referencia."])
    center_validation = DataValidation(
        type="list", formula1=f"'Lineas de negocio'!$A$2:$A${max(2, len(centers) + 1)}", allow_blank=True
    )
    rubro_validation = DataValidation(
        type="list", formula1=f"'Rubros'!$A$2:$A${max(2, len(rubros) + 1)}", allow_blank=False
    )
    for validation, target in (
        (rubro_validation, "B2:B5000"),
        (center_validation, "D2:D5000"),
    ):
        sheet.add_data_validation(validation)
        validation.add(target)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_egyp.xlsx"},
    )


def resolve_row(
    db: Session, values: dict[str, Any], row_number: int
) -> tuple[RubroResultado, CentroResultado | None]:
    if not values["rubro_codigo"]:
        raise ValueError(f"Fila {row_number}: rubro_codigo es obligatorio.")
    rubro = db.scalar(
        select(RubroResultado).where(
            RubroResultado.codigo == values["rubro_codigo"], RubroResultado.is_active.is_(True)
        )
    )
    centro = None
    if values["centro_codigo"]:
        centro = db.scalar(
            select(CentroResultado).where(
                CentroResultado.codigo == values["centro_codigo"],
                CentroResultado.is_active.is_(True),
            )
        )
    if not rubro:
        raise ValueError(f"Fila {row_number}: rubro_codigo no existe o esta inactivo.")
    if values["centro_codigo"] and not centro:
        raise ValueError(f"Fila {row_number}: centro_codigo no existe o esta inactivo.")
    if not centro and not active_distribution_rule(
        db, values["fecha"], rubro.id
    ):
        raise ValueError(
            f"Fila {row_number}: centro_codigo es obligatorio si no existe una regla de distribucion vigente."
        )
    return rubro, centro


def active_distribution_rule(
    db: Session, entry_date: date, rubro_id: uuid.UUID
) -> ReglaDistribucion | None:
    rules = list(
        db.scalars(
            select(ReglaDistribucion)
            .where(
                ReglaDistribucion.is_active.is_(True),
                (ReglaDistribucion.vigente_desde.is_(None))
                | (ReglaDistribucion.vigente_desde <= entry_date),
                (ReglaDistribucion.vigente_hasta.is_(None))
                | (ReglaDistribucion.vigente_hasta >= entry_date),
            )
            .order_by(ReglaDistribucion.nombre)
        )
    )
    return next((rule for rule in rules if rule.rubro_id == rubro_id), None)


def create_entry(
    db: Session,
    user: User,
    values: dict[str, Any],
    rubro: RubroResultado,
    centro: CentroResultado | None,
    origen: OrigenAsientoResultado,
    lote_id: uuid.UUID | None = None,
    numero_fila: int | None = None,
) -> AsientoResultado:
    entry_values = {
        key: values[key]
        for key in (
            "fecha",
            "descripcion",
            "documento",
            "moneda",
            "debe",
            "haber",
            "observaciones",
        )
    }
    if centro:
        entry = AsientoResultado(
            **entry_values,
            rubro_id=rubro.id,
            centro_resultado_id=centro.id,
            origen=origen,
            lote_id=lote_id,
            numero_fila=numero_fila,
            created_by_id=user.id,
            updated_by_id=user.id,
        )
        db.add(entry)
        return entry
    rule = active_distribution_rule(db, values["fecha"], rubro.id)
    if not rule:
        entry = AsientoResultado(
            **entry_values,
            rubro_id=rubro.id,
            origen=origen,
            lote_id=lote_id,
            numero_fila=numero_fila,
            created_by_id=user.id,
            updated_by_id=user.id,
        )
        db.add(entry)
        return entry
    origin_entry = AsientoResultado(
        **entry_values,
        rubro_id=rubro.id,
        origen=origen,
        es_resultado=False,
        lote_id=lote_id,
        numero_fila=numero_fila,
        created_by_id=user.id,
        updated_by_id=user.id,
    )
    db.add(origin_entry)
    db.flush()
    remaining_debe, remaining_haber = Decimal(values["debe"]), Decimal(values["haber"])
    for index, line in enumerate(rule.lineas):
        is_last = index == len(rule.lineas) - 1
        debe = (
            remaining_debe
            if is_last
            else (Decimal(values["debe"]) * Decimal(line.porcentaje) / 100).quantize(
                Decimal("0.01")
            )
        )
        haber = (
            remaining_haber
            if is_last
            else (Decimal(values["haber"]) * Decimal(line.porcentaje) / 100).quantize(
                Decimal("0.01")
            )
        )
        remaining_debe -= debe
        remaining_haber -= haber
        db.add(
            AsientoResultado(
                **{**entry_values, "debe": debe, "haber": haber},
                rubro_id=rubro.id,
                centro_resultado_id=line.centro_resultado_id,
                origen=OrigenAsientoResultado.DISTRIBUIDO,
                asiento_origen_id=origin_entry.id,
                lote_id=lote_id,
                created_by_id=user.id,
                updated_by_id=user.id,
            )
        )
    return origin_entry


@router.post(
    "/importar.xlsx",
    response_model=ImportResultadoRead,
    status_code=status.HTTP_201_CREATED,
    dependencies=[Depends(require_enabled)],
)
async def import_asientos(
    archivo: Annotated[UploadFile, File(...)],
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    if not (archivo.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="El archivo debe ser .xlsx."
        )
    try:
        workbook = load_workbook(BytesIO(await archivo.read()), read_only=True, data_only=True)
        sheet = workbook["Asientos"]
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="No se pudo leer la hoja Asientos.",
        ) from exc
    headers = [
        str(value or "").strip()
        for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    ]
    if headers != TEMPLATE_HEADERS:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Las cabeceras de la hoja Asientos no coinciden con la plantilla.",
        )
    parsed: list[dict[str, Any]] = []
    errors: list[dict[str, object]] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        try:
            raw = dict(zip(TEMPLATE_HEADERS, row, strict=True))
            period = str(raw["periodo"] or "").strip()
            if len(period) != 7 or period[4] != "-":
                raise ValueError("periodo debe tener formato AAAA-MM.")
            year, month = (int(value) for value in period.split("-"))
            entry_date = date(year, month, monthrange(year, month)[1])
            amount = Decimal(str(raw["monto_sin_igv"] or 0))
            if amount <= ZERO:
                raise ValueError("monto_sin_igv debe ser mayor que cero.")
            values = {
                "fecha": entry_date,
                "centro_codigo": str(raw["centro_codigo"] or "").strip(),
                "rubro_codigo": str(raw["rubro_codigo"] or "").strip(),
                "numero_fila": row_number,
            }
            rubro, _ = resolve_row(db, values, row_number)
            credit = rubro.naturaleza in {
                NaturalezaRubro.INGRESO,
                NaturalezaRubro.INGRESO_FINANCIERO,
            }
            parsed.append(
                {
                    **values,
                    "descripcion": str(raw["descripcion"] or "").strip() or rubro.nombre,
                    "documento": f"Cierre {period}",
                    "moneda": "PEN",
                    "debe": ZERO if credit else amount,
                    "haber": amount if credit else ZERO,
                    "observaciones": str(raw["observaciones"] or "").strip() or None,
                }
            )
        except (ValueError, TypeError, ArithmeticError) as exc:
            errors.append({"fila": row_number, "mensaje": str(exc)})
    if errors:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail={"errores": errors}
        )
    batch = LoteResultado(
        archivo_nombre=archivo.filename or "asientos.xlsx",
        usuario_id=user.id,
        total_filas=len(parsed),
    )
    db.add(batch)
    db.flush()
    for row in parsed:
        rubro, centro = resolve_row(db, row, row["numero_fila"])
        create_entry(
            db,
            user,
            row,
            rubro,
            centro,
            OrigenAsientoResultado.IMPORTADO,
            batch.id,
            row["numero_fila"],
        )
    log_audit(
        db,
        user=user,
        table=batch.__tablename__,
        record_id=batch.id,
        action=AuditAction.CREATE,
        changes={"archivo_nombre": batch.archivo_nombre, "total_filas": len(parsed)},
    )
    db.commit()
    return ImportResultadoRead(lote_id=batch.id, total_filas=len(parsed), filas_nuevas=len(parsed))


@router.post(
    "/asientos",
    response_model=AsientoResultadoRead,
    status_code=status.HTTP_201_CREATED,
    dependencies=[Depends(require_enabled)],
)
def create_manual_entry(
    data: AsientoResultadoCreate,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    rubro = require_active(db, RubroResultado, data.rubro_id, "Rubro")
    centro = (
        require_active(db, CentroResultado, data.centro_resultado_id, "Linea de negocio")
        if data.centro_resultado_id
        else None
    )
    if not centro and not active_distribution_rule(db, data.fecha, rubro.id):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="La linea de negocio es obligatoria si no existe una regla de distribucion vigente.",
        )
    record = create_entry(db, user, data.model_dump(), rubro, centro, OrigenAsientoResultado.MANUAL)
    db.flush()
    log_audit(
        db,
        user=user,
        table=record.__tablename__,
        record_id=record.id,
        action=AuditAction.CREATE,
        changes={"origen": OrigenAsientoResultado.MANUAL.value, **data.model_dump(mode="json")},
    )
    db.commit()
    db.refresh(record)
    return record


@router.patch(
    "/asientos/{asiento_id}",
    response_model=AsientoResultadoRead,
    dependencies=[Depends(require_enabled)],
)
def update_manual_entry(
    asiento_id: uuid.UUID,
    data: AsientoResultadoUpdate,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    record = require_active(db, AsientoResultado, asiento_id, "Asiento")
    if record.origen != OrigenAsientoResultado.MANUAL or not record.es_resultado:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Solo se pueden editar asientos manuales no distribuidos.",
        )
    values = data.model_dump(exclude_unset=True)
    merged = {
        field: values.get(field, getattr(record, field))
        for field in AsientoResultadoCreate.model_fields
    }
    if merged["debe"] == ZERO and merged["haber"] == ZERO:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Debe u haber es obligatorio."
        )
    rubro = require_active(db, RubroResultado, merged["rubro_id"], "Rubro")
    if merged["centro_resultado_id"]:
        require_active(db, CentroResultado, merged["centro_resultado_id"], "Linea de negocio")
    elif not active_distribution_rule(db, merged["fecha"], rubro.id):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="La linea de negocio es obligatoria si no existe una regla de distribucion vigente.",
        )
    changes = apply_update(record, values)
    if changes:
        record.updated_by_id = user.id
        save_change(db, user, record, AuditAction.UPDATE, changes)
    return record


@router.patch(
    "/asientos/{asiento_id}/anular",
    response_model=AsientoResultadoRead,
    dependencies=[Depends(require_enabled)],
)
def cancel_entry(
    asiento_id: uuid.UUID,
    data: AnularAsientoResultado,
    user: Annotated[User, Depends(require_admin)],
    db: Annotated[Session, Depends(get_db)],
):
    record = require_active(db, AsientoResultado, asiento_id, "Asiento")
    entries = list(
        db.scalars(
            select(AsientoResultado).where(
                or_(
                    AsientoResultado.id == asiento_id,
                    AsientoResultado.asiento_origen_id == asiento_id,
                ),
                AsientoResultado.is_active.is_(True),
            )
        )
    )
    for entry in entries:
        entry.is_active = False
        entry.deleted_at = datetime.now(UTC)
        entry.motivo_anulacion = data.motivo
        entry.anulado_by_id = user.id
        entry.updated_by_id = user.id
    log_audit(
        db,
        user=user,
        table=record.__tablename__,
        record_id=record.id,
        action=AuditAction.SOFT_DELETE,
        changes={
            "motivo_anulacion": data.motivo,
            "asientos_anulados": [str(entry.id) for entry in entries],
        },
    )
    db.commit()
    db.refresh(record)
    return record


@router.get(
    "/asientos", response_model=list[AsientoResultadoRead], dependencies=[Depends(require_enabled)]
)
def list_asientos(
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    rubro_id: uuid.UUID | None = None,
    centro_id: uuid.UUID | None = None,
    _: Annotated[User, Depends(get_current_user)] = None,
    db: Annotated[Session, Depends(get_db)] = None,
):
    statement = select(AsientoResultado).where(AsientoResultado.is_active.is_(True))
    if fecha_desde:
        statement = statement.where(AsientoResultado.fecha >= fecha_desde)
    if fecha_hasta:
        statement = statement.where(AsientoResultado.fecha <= fecha_hasta)
    if rubro_id:
        statement = statement.where(AsientoResultado.rubro_id == rubro_id)
    if centro_id:
        statement = statement.where(AsientoResultado.centro_resultado_id == centro_id)
    return list(
        db.scalars(
            statement.order_by(AsientoResultado.fecha.desc(), AsientoResultado.created_at.desc())
        )
    )


def signed_amount(entry: AsientoResultado) -> Decimal:
    # Credits increase income; debits increase costs and expenses, which reduce the result.
    return Decimal(entry.haber) - Decimal(entry.debe)


@router.get(
    "/resumen", response_model=ResumenGananciasPerdidasRead, dependencies=[Depends(require_enabled)]
)
def summary(
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    _: Annotated[User, Depends(get_current_user)] = None,
    db: Annotated[Session, Depends(get_db)] = None,
):
    if fecha_desde and fecha_hasta and fecha_desde > fecha_hasta:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="fecha_desde no puede ser posterior a fecha_hasta.",
        )
    all_centers = list(
        db.scalars(
            select(CentroResultado)
            .where(CentroResultado.is_active.is_(True))
            .order_by(CentroResultado.orden, CentroResultado.nombre)
        )
    )
    rubros = list(
        db.scalars(
            select(RubroResultado)
            .where(RubroResultado.is_active.is_(True))
            .order_by(RubroResultado.orden, RubroResultado.nombre)
        )
    )
    by_id = {rubro.id: rubro for rubro in rubros}
    totals = {
        rubro.id: {"total": ZERO, "centers": {center.codigo: ZERO for center in all_centers}}
        for rubro in rubros
    }
    total_by_nature = {nature: ZERO for nature in NaturalezaRubro}
    totals_by_center = {
        center.id: {nature: ZERO for nature in NaturalezaRubro}
        for center in all_centers
    }
    statement = select(AsientoResultado).where(
        AsientoResultado.is_active.is_(True), AsientoResultado.es_resultado.is_(True)
    )
    if fecha_desde:
        statement = statement.where(AsientoResultado.fecha >= fecha_desde)
    if fecha_hasta:
        statement = statement.where(AsientoResultado.fecha <= fecha_hasta)
    center_ids = {center.id: center.codigo for center in all_centers}
    used_center_ids: set[uuid.UUID] = set()
    for entry in db.scalars(statement):
        rubro = by_id.get(entry.rubro_id)
        if not rubro:
            continue
        amount = signed_amount(entry)
        total_by_nature[rubro.naturaleza] += amount
        if entry.centro_resultado_id in totals_by_center:
            totals_by_center[entry.centro_resultado_id][rubro.naturaleza] += amount
            used_center_ids.add(entry.centro_resultado_id)
        current: RubroResultado | None = rubro
        while current:
            totals[current.id]["total"] += amount
            if entry.centro_resultado_id in center_ids:
                totals[current.id]["centers"][center_ids[entry.centro_resultado_id]] += amount
            current = by_id.get(current.padre_id) if current.padre_id else None
    centers = [center for center in all_centers if center.id in used_center_ids]
    result_rows = [
        ResultadoRubroRead(
            rubro_id=rubro.id,
            codigo=rubro.codigo,
            nombre=rubro.nombre,
            naturaleza=rubro.naturaleza,
            padre_id=rubro.padre_id,
            total=totals[rubro.id]["total"],
            por_centro=totals[rubro.id]["centers"],
        )
        for rubro in rubros
    ]
    ingresos_brutos = total_by_nature[NaturalezaRubro.INGRESO]
    ingresos = (
        total_by_nature[NaturalezaRubro.INGRESO]
        + total_by_nature[NaturalezaRubro.INGRESO_FINANCIERO]
    )
    ingresos_netos = (
        total_by_nature[NaturalezaRubro.INGRESO] + total_by_nature[NaturalezaRubro.CONTRA_INGRESO]
    )
    utilidad_bruta = ingresos_netos + total_by_nature[NaturalezaRubro.COSTO_VENTA]
    gastos_operativos = total_by_nature[NaturalezaRubro.GASTO_OPERATIVO]
    utilidad_operativa = utilidad_bruta + gastos_operativos
    rentabilidad_por_centro = []
    for center in centers:
        values = totals_by_center[center.id]
        ingresos_netos_centro = values[NaturalezaRubro.INGRESO] + values[NaturalezaRubro.CONTRA_INGRESO]
        utilidad_bruta_centro = ingresos_netos_centro + values[NaturalezaRubro.COSTO_VENTA]
        utilidad_neta_centro = (
            values[NaturalezaRubro.INGRESO]
            + values[NaturalezaRubro.INGRESO_FINANCIERO]
            + values[NaturalezaRubro.CONTRA_INGRESO]
            + values[NaturalezaRubro.COSTO_VENTA]
            + values[NaturalezaRubro.GASTO_OPERATIVO]
            + values[NaturalezaRubro.GASTO_FINANCIERO]
            + values[NaturalezaRubro.IMPUESTO]
        )
        rentabilidad_por_centro.append(RentabilidadCentroResultadoRead(
            centro_id=center.id,
            codigo=center.codigo,
            nombre=center.nombre,
            ingresos_brutos=values[NaturalezaRubro.INGRESO],
            ingresos_netos=ingresos_netos_centro,
            utilidad_bruta=utilidad_bruta_centro,
            utilidad_neta=utilidad_neta_centro,
        ))
    return ResumenGananciasPerdidasRead(
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        centros=centers,
        rubros=result_rows,
        ingresos_brutos=ingresos_brutos,
        ingresos_netos=ingresos_netos,
        utilidad_bruta=utilidad_bruta,
        gastos_operativos=gastos_operativos,
        utilidad_operativa=utilidad_operativa,
        utilidad_neta=ingresos
        + total_by_nature[NaturalezaRubro.CONTRA_INGRESO]
        + total_by_nature[NaturalezaRubro.COSTO_VENTA]
        + gastos_operativos
        + total_by_nature[NaturalezaRubro.GASTO_FINANCIERO]
        + total_by_nature[NaturalezaRubro.IMPUESTO],
        rentabilidad_por_centro=rentabilidad_por_centro,
    )
