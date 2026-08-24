import io
import re
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Annotated, Any
from zipfile import BadZipFile

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.encoders import jsonable_encoder
from fastapi.responses import JSONResponse, Response
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, aliased

from app.core.database import get_db
from app.modules.categorias.models import Actividad, Concepto, Tipo
from app.modules.cuentas_bancos.models import Banco, CuentaBancaria
from app.modules.reports.schemas import ActualizacionMasivaResponse
from app.modules.shared.audit import log_audit
from app.modules.shared.models import AuditAction
from app.modules.sucursales.models import Sucursal
from app.modules.transacciones.models import EstadoTransaccion, OrigenTransaccion, Transaccion
from app.modules.usuarios.deps import get_current_user
from app.modules.usuarios.models import User, UserRole
from app.modules.vendedores.models import Vendedor

router = APIRouter(prefix="/reports", tags=["reports"])

HEADERS = [
    "fecha",
    "descripcion",
    "n_operacion",
    "monto",
    "documento",
    "sucursal",
    "vendedor",
    "observaciones",
    "actividad",
    "concepto",
    "tipo",
    "fuente",
]

AUDIT_HEADERS = [
    "fecha",
    "descripcion",
    "n_operacion",
    "monto",
    "moneda",
    "documento",
    "sucursal",
    "vendedor",
    "observaciones",
    "actividad",
    "concepto",
    "tipo",
    "fuente",
    "origen",
    "estado",
    "es_proyeccion",
    "materializado",
    "activo",
    "registrado_por",
    "actualizado_por",
    "creado_el",
    "actualizado_el",
]


def fuente(banco_codigo: str, cuenta_alias: str) -> str:
    return f"{banco_codigo} | {cuenta_alias}"


def response_payload(
    confirmar: bool,
    total_filas: int,
    errores: list[dict[str, object]],
    cambios: list[dict[str, object]],
) -> dict[str, object]:
    return {
        "confirmar": confirmar,
        "aplicado": False,
        "total_filas": total_filas,
        "filas_validas": total_filas - len(errores),
        "filas_error": len(errores),
        "errores": errores,
        "cambios": cambios,
    }


def cell_value(value: object) -> str:
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value or "").strip()


def read_rows(content: bytes) -> tuple[list[dict[str, str]], list[dict[str, object]], int]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError):
        return [], [{"fila": 0, "mensaje": "El archivo debe ser un Excel (.xlsx) valido."}], 0

    worksheet = workbook.worksheets[0]
    raw_rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if not raw_rows or [cell_value(value) for value in raw_rows[0]] != HEADERS:
        return [], [{"fila": 1, "mensaje": "Las cabeceras no son validas."}], 0
    raw_rows = raw_rows[1:]
    if not raw_rows:
        return [], [{"fila": 2, "mensaje": "El archivo no contiene movimientos."}], 0

    rows: list[dict[str, str]] = []
    errors: list[dict[str, object]] = []
    for row_number, row in enumerate(raw_rows, start=2):
        values = [cell_value(value) for value in row]
        if len(values) != len(HEADERS) or all(not value for value in values):
            errors.append(
                {"fila": row_number, "mensaje": "No se permiten filas vacias o columnas extra."}
            )
            continue
        rows.append({"fila": str(row_number)} | dict(zip(HEADERS, values, strict=True)))
    return rows, errors, len(raw_rows)


def active_by_name(db: Session, model: Any, name: str) -> Any | None:
    records = list(db.scalars(select(model).where(model.nombre == name, model.is_active.is_(True))))
    if len(records) != 1:
        return None
    return records[0]


def parse_date(value: str) -> date | None:
    cleaned = value.strip().replace("\xa0", "")
    if cleaned.startswith("="):
        cleaned = cleaned[1:].strip()
    cleaned = cleaned.strip('"').strip("'").strip().split("T", 1)[0]
    if not cleaned:
        return None
    match = re.search(r"\d{1,4}[/.\-]\d{1,2}[/.\-]\d{1,4}", cleaned)
    if match:
        cleaned = match.group()
    else:
        cleaned = cleaned.split(maxsplit=1)[0]
    for fmt in (
        "%d/%m/%Y",
        "%m/%d/%y",
        "%d-%m-%Y",
        "%m-%d-%y",
        "%d.%m.%Y",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
    ):
        try:
            return datetime.strptime(cleaned, fmt).date()
        except ValueError:
            continue
    return None


def validate_row(
    db: Session,
    row: dict[str, str],
    transactions: dict[tuple[str, str], list[Transaccion]],
) -> tuple[dict[str, Any] | None, str | None]:
    parsed_date = parse_date(row["fecha"])
    if not parsed_date:
        return None, "fecha debe tener formato DD/MM/AAAA (ej: 15/08/2026)."
    if not row["descripcion"]:
        return None, "descripcion es obligatoria."
    if not row["n_operacion"] or not row["fuente"]:
        return None, "n_operacion y fuente son obligatorios."
    try:
        monto = Decimal(row["monto"].replace(",", "."))
    except (InvalidOperation, ValueError):
        return None, "monto debe ser un decimal valido."
    if not monto.is_finite() or monto == 0:
        return None, "monto debe ser finito y distinto de cero."

    matched = transactions.get((row["n_operacion"], row["fuente"]), [])
    if len(matched) != 1:
        return None, "No se encontro una transaccion activa para n_operacion y fuente."
    references: dict[str, Any | None] = {}
    for field, model, label in (
        ("sucursal", Sucursal, "Sucursal"),
        ("vendedor", Vendedor, "Vendedor"),
        ("actividad", Actividad, "Actividad"),
    ):
        if row[field]:
            reference = active_by_name(db, model, row[field])
            if reference is None:
                return None, f"{label} activo no encontrado o duplicado: {row[field]}."
            references[field] = reference
        else:
            references[field] = None

    actividad_ref = references.get("actividad")
    if row.get("concepto"):
        if not actividad_ref:
            return None, "concepto requiere que actividad tambien este indicada."
        concepto_ref = db.scalar(
            select(Concepto).where(
                Concepto.nombre == row["concepto"],
                Concepto.actividad_id == actividad_ref.id,
                Concepto.is_active.is_(True),
            )
        )
        if not concepto_ref:
            return None, (
                f"Concepto activo no encontrado para la actividad "
                f"'{row['actividad']}': {row['concepto']}."
            )
        references["concepto"] = concepto_ref
    else:
        references["concepto"] = None

    concepto_ref = references.get("concepto")
    if row.get("tipo"):
        if not concepto_ref:
            return None, "tipo requiere que concepto tambien este indicada."
        tipo_ref = db.scalar(
            select(Tipo).where(
                Tipo.nombre == row["tipo"],
                Tipo.concepto_id == concepto_ref.id,
                Tipo.is_active.is_(True),
            )
        )
        if not tipo_ref:
            return None, (
                f"Tipo activo no encontrado para el concepto "
                f"'{row['concepto']}': {row['tipo']}."
            )
        references["tipo"] = tipo_ref
    else:
        references["tipo"] = None

    categories = (references["actividad"], references["concepto"], references["tipo"])
    if any(categories) and not all(categories):
        return None, "actividad, concepto y tipo deben estar todos completos o vacios."
    return {
        "transaction": matched[0],
        "fecha": parsed_date,
        "descripcion": row["descripcion"],
        "monto": monto,
        "documento": row["documento"] or None,
        "observaciones": row["observaciones"] or None,
        "sucursal_id": references["sucursal"].id if references["sucursal"] else None,
        "vendedor_id": references["vendedor"].id if references["vendedor"] else None,
        "actividad_id": references["actividad"].id if references["actividad"] else None,
        "concepto_id": references["concepto"].id if references["concepto"] else None,
        "tipo_id": references["tipo"].id if references["tipo"] else None,
    }, None


@router.get("/no-tipificados.xlsx")
def export_no_tipificados(
    created_by_id: uuid.UUID | None = Query(default=None),
    current_user: Annotated[User, Depends(get_current_user)] = None,
    db: Annotated[Session, Depends(get_db)] = None,
) -> Response:
    if created_by_id and current_user.rol != UserRole.ADMIN and created_by_id != current_user.id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Solo puedes descargar tus propios pendientes.",
        )
    filters = [
        Transaccion.is_active.is_(True),
        Transaccion.estado == EstadoTransaccion.PENDIENTE_TIPIFICAR,
        Transaccion.n_operacion.is_not(None),
    ]
    if created_by_id:
        filters.append(Transaccion.created_by_id == created_by_id)
    statement = (
        select(
            Transaccion,
            Banco.codigo,
            CuentaBancaria.alias,
            Sucursal.nombre,
            Vendedor.nombre,
            Actividad.nombre,
            Concepto.nombre,
            Tipo.nombre,
        )
        .join(CuentaBancaria, Transaccion.cuenta_bancaria_id == CuentaBancaria.id)
        .join(Banco, CuentaBancaria.banco_id == Banco.id)
        .outerjoin(Sucursal, Transaccion.sucursal_id == Sucursal.id)
        .outerjoin(Vendedor, Transaccion.vendedor_id == Vendedor.id)
        .outerjoin(Actividad, Transaccion.actividad_id == Actividad.id)
        .outerjoin(Concepto, Transaccion.concepto_id == Concepto.id)
        .outerjoin(Tipo, Transaccion.tipo_id == Tipo.id)
        .where(*filters)
        .order_by(Transaccion.fecha, Transaccion.id)
    )
    workbook = Workbook()
    movements = workbook.active
    movements.title = "Movimientos"
    movements.append(HEADERS)
    for (
        transaction,
        banco_codigo,
        alias,
        sucursal,
        vendedor,
        actividad,
        concepto,
        tipo,
    ) in db.execute(statement):
        movements.append(
            [
                transaction.fecha.strftime("%d/%m/%Y"),
                transaction.descripcion,
                transaction.n_operacion,
                transaction.monto,
                transaction.documento or "",
                sucursal or "",
                vendedor or "",
                transaction.observaciones or "",
                actividad or "",
                concepto or "",
                tipo or "",
                fuente(banco_codigo, alias),
            ]
        )
    tipifications = workbook.create_sheet("Tipificaciones")
    tipifications.append(["actividad", "concepto", "tipo"])
    for actividad, concepto, tipo in db.execute(
        select(Actividad.nombre, Concepto.nombre, Tipo.nombre)
        .join(Concepto, Tipo.concepto_id == Concepto.id)
        .join(Actividad, Concepto.actividad_id == Actividad.id)
        .where(
            Actividad.is_active.is_(True),
            Concepto.is_active.is_(True),
            Tipo.is_active.is_(True),
        )
        .order_by(Actividad.nombre, Concepto.nombre, Tipo.nombre)
    ):
        tipifications.append([actividad, concepto, tipo])
    output = io.BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="no-tipificados.xlsx"'},
    )


@router.get("/auditoria.xlsx")
def export_auditoria(
    cuenta_bancaria_id: uuid.UUID | None = None,
    actividad_id: uuid.UUID | None = None,
    concepto_id: uuid.UUID | None = None,
    tipo_id: uuid.UUID | None = None,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    tipo_movimiento: str = "TODOS",
    estado_registro: str = "TODOS",
    estado: EstadoTransaccion | None = None,
    _: Annotated[User, Depends(get_current_user)] = None,
    db: Annotated[Session, Depends(get_db)] = None,
) -> Response:
    if fecha_desde and fecha_hasta and fecha_desde > fecha_hasta:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="fecha_desde no puede ser posterior a fecha_hasta.",
        )
    if tipo_movimiento not in {"TODOS", "PROYECCIONES", "REALES"}:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="tipo_movimiento debe ser TODOS, PROYECCIONES o REALES.",
        )
    if estado_registro not in {"TODOS", "ACTIVOS", "ANULADOS"}:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="estado_registro debe ser TODOS, ACTIVOS o ANULADOS.",
        )
    filters = []
    if estado_registro == "ACTIVOS":
        filters.append(Transaccion.is_active.is_(True))
    elif estado_registro == "ANULADOS":
        filters.append(Transaccion.is_active.is_(False))
    if tipo_movimiento == "PROYECCIONES":
        filters.append(Transaccion.origen == OrigenTransaccion.PROYECCION)
    elif tipo_movimiento == "REALES":
        filters.append(
            Transaccion.origen.in_(
                [OrigenTransaccion.IMPORTADO, OrigenTransaccion.MANUAL]
            )
        )
    if cuenta_bancaria_id:
        filters.append(Transaccion.cuenta_bancaria_id == cuenta_bancaria_id)
    if actividad_id:
        filters.append(Transaccion.actividad_id == actividad_id)
    if concepto_id:
        filters.append(Transaccion.concepto_id == concepto_id)
    if tipo_id:
        filters.append(Transaccion.tipo_id == tipo_id)
    if estado:
        filters.append(Transaccion.estado == estado)
    if fecha_desde:
        filters.append(Transaccion.fecha >= fecha_desde)
    if fecha_hasta:
        filters.append(Transaccion.fecha <= fecha_hasta)

    creator = aliased(User)
    updater = aliased(User)
    statement = (
        select(
            Transaccion,
            Banco.codigo,
            CuentaBancaria.alias,
            Sucursal.nombre,
            Vendedor.nombre,
            Actividad.nombre,
            Concepto.nombre,
            Tipo.nombre,
            creator.nombre,
            updater.nombre,
        )
        .join(CuentaBancaria, Transaccion.cuenta_bancaria_id == CuentaBancaria.id)
        .join(Banco, CuentaBancaria.banco_id == Banco.id)
        .outerjoin(Sucursal, Transaccion.sucursal_id == Sucursal.id)
        .outerjoin(Vendedor, Transaccion.vendedor_id == Vendedor.id)
        .outerjoin(Actividad, Transaccion.actividad_id == Actividad.id)
        .outerjoin(Concepto, Transaccion.concepto_id == Concepto.id)
        .outerjoin(Tipo, Transaccion.tipo_id == Tipo.id)
        .join(creator, Transaccion.created_by_id == creator.id)
        .join(updater, Transaccion.updated_by_id == updater.id)
        .where(*filters)
        .order_by(Transaccion.fecha, Transaccion.created_at, Transaccion.id)
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Auditoria"
    sheet.append(AUDIT_HEADERS)
    for row in db.execute(statement):
        (
            transaction,
            banco_codigo,
            alias,
            sucursal,
            vendedor,
            actividad,
            concepto,
            tipo,
            creator_name,
            updater_name,
        ) = row
        sheet.append(
            [
                transaction.fecha.strftime("%d/%m/%Y"),
                transaction.descripcion,
                transaction.n_operacion or "",
                transaction.monto,
                transaction.moneda,
                transaction.documento or "",
                sucursal or "",
                vendedor or "",
                transaction.observaciones or "",
                actividad or "",
                concepto or "",
                tipo or "",
                fuente(banco_codigo, alias),
                transaction.origen.value,
                transaction.estado.value,
                transaction.origen == OrigenTransaccion.PROYECCION,
                transaction.materializado,
                transaction.is_active,
                creator_name,
                updater_name,
                transaction.created_at.isoformat(),
                transaction.updated_at.isoformat(),
            ]
        )
    output = io.BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="auditoria_transacciones.xlsx"'},
    )


@router.post("/actualizacion-masiva", response_model=ActualizacionMasivaResponse)
def actualizacion_masiva(
    archivo: Annotated[UploadFile, File()],
    confirmar: Annotated[bool, Form()] = False,
    current_user: Annotated[User, Depends(get_current_user)] = None,
    db: Annotated[Session, Depends(get_db)] = None,
) -> ActualizacionMasivaResponse | JSONResponse:
    rows, errors, total_filas = read_rows(archivo.file.read())
    if errors and not rows:
        payload = response_payload(confirmar, total_filas, errors, [])
        return (
            JSONResponse(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                content=jsonable_encoder(payload),
            )
            if confirmar
            else payload
        )

    transactions: dict[tuple[str, str], list[Transaccion]] = {}
    for transaction, codigo, alias in db.execute(
        select(Transaccion, Banco.codigo, CuentaBancaria.alias)
        .join(CuentaBancaria, Transaccion.cuenta_bancaria_id == CuentaBancaria.id)
        .join(Banco, CuentaBancaria.banco_id == Banco.id)
        .where(Transaccion.is_active.is_(True), Transaccion.n_operacion.is_not(None))
    ):
        transactions.setdefault((transaction.n_operacion, fuente(codigo, alias)), []).append(
            transaction
        )

    seen: set[tuple[str, str]] = set()
    updates: list[tuple[int, dict[str, Any]]] = []
    for row in rows:
        row_number = int(row["fila"])
        key = (row["n_operacion"], row["fuente"])
        if key in seen:
            errors.append(
                {"fila": row_number, "mensaje": "Fila duplicada para n_operacion y fuente."}
            )
            continue
        seen.add(key)
        values, message = validate_row(db, row, transactions)
        if message:
            errors.append({"fila": row_number, "mensaje": message})
        else:
            updates.append((row_number, values))

    changes_response: list[dict[str, object]] = []
    for row_number, values in updates:
        transaction = values["transaction"]
        changes = {
            field: {"old": getattr(transaction, field), "new": values[field]}
            for field in values
            if field != "transaction" and getattr(transaction, field) != values[field]
        }
        complete = values["actividad_id"] is not None
        next_state = None
        if transaction.origen != OrigenTransaccion.PROYECCION:
            if complete:
                next_state = EstadoTransaccion.TIPIFICADO
            elif transaction.origen == OrigenTransaccion.IMPORTADO:
                next_state = EstadoTransaccion.PENDIENTE_TIPIFICAR
        if next_state is not None and transaction.estado != next_state:
            changes["estado"] = {"old": transaction.estado, "new": next_state}
        changes_response.append(
            {"fila": row_number, "transaccion_id": transaction.id, "cambios": changes}
        )

    payload = response_payload(confirmar, total_filas, errors, changes_response)
    if not confirmar:
        return payload
    if errors:
        return JSONResponse(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            content=jsonable_encoder(payload),
        )
    try:
        for _, values in updates:
            transaction = values["transaction"]
            changes = next(
                item["cambios"]
                for item in changes_response
                if item["transaccion_id"] == transaction.id
            )
            if not changes:
                continue
            for field, value in values.items():
                if field != "transaction":
                    setattr(transaction, field, value)
            if "estado" in changes:
                transaction.estado = changes["estado"]["new"]
            transaction.updated_by_id = current_user.id
            log_audit(
                db,
                user=current_user,
                table=Transaccion.__tablename__,
                record_id=transaction.id,
                action=AuditAction.UPDATE,
                changes=changes,
            )
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT, detail="Conflicto al actualizar transacciones."
        ) from exc
    payload["aplicado"] = True
    return payload
