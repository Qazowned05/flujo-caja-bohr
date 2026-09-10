from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation


def add_dependent_tipification_lists(
    workbook: Workbook,
    movements,
    *,
    tipification_rows: Sequence[tuple[str, str, str]],
    branch_names: Sequence[str],
    seller_rows: Sequence[tuple[str, str]],
    activity_column: str,
    concept_column: str,
    type_column: str,
    branch_column: str | None = None,
    seller_column: str | None = None,
    allow_blank: bool,
) -> None:
    lists = workbook.create_sheet("Listas tipificacion")
    lists.append(["actividad", "rango_conceptos", "", "actividad_concepto", "rango_tipos", "", "actividades", "", "sucursal", "rango_vendedores", "", "vacio"])
    lists.append(["", "", "", "", "", "", "", "", "", "", "", ""])
    workbook.defined_names.add(DefinedName("empty_list", attr_text="'Listas tipificacion'!$L$2"))

    activity_groups: dict[str, dict[str, list[str]]] = {}
    for activity, concept, kind in tipification_rows:
        activity_groups.setdefault(activity, {}).setdefault(concept, []).append(kind)
    activity_names = sorted(activity_groups)
    if activity_names:
        for row, activity in enumerate(activity_names, start=2):
            lists.cell(row=row, column=1, value=activity)
            lists.cell(row=row, column=7, value=activity)
        workbook.defined_names.add(
            DefinedName("activity_options", attr_text=f"'Listas tipificacion'!$G$2:$G${len(activity_names) + 1}")
        )
    else:
        workbook.defined_names.add(DefinedName("activity_options", attr_text="'Listas tipificacion'!$L$2"))

    list_column = 13
    type_range_number = 0
    for activity_number, activity in enumerate(activity_names, start=1):
        concepts = sorted(activity_groups[activity])
        concept_name = f"concept_options_{activity_number}"
        column_letter = get_column_letter(list_column)
        for row, concept in enumerate(concepts, start=2):
            lists.cell(row=row, column=list_column, value=concept)
        workbook.defined_names.add(
            DefinedName(concept_name, attr_text=f"'Listas tipificacion'!${column_letter}$2:${column_letter}${max(2, len(concepts) + 1)}")
        )
        lists.cell(row=activity_number + 1, column=2, value=concept_name)
        list_column += 1
        for concept in concepts:
            type_range_number += 1
            type_name = f"type_options_{type_range_number}"
            kinds = sorted(set(activity_groups[activity][concept]))
            column_letter = get_column_letter(list_column)
            for row, kind in enumerate(kinds, start=2):
                lists.cell(row=row, column=list_column, value=kind)
            workbook.defined_names.add(
                DefinedName(type_name, attr_text=f"'Listas tipificacion'!${column_letter}$2:${column_letter}${max(2, len(kinds) + 1)}")
            )
            lists.cell(row=type_range_number + 1, column=4, value=f"{activity}|{concept}")
            lists.cell(row=type_range_number + 1, column=5, value=type_name)
            list_column += 1

    seller_groups: dict[str, list[str]] = {branch: [] for branch in branch_names}
    for seller, branch in seller_rows:
        if branch in seller_groups:
            seller_groups[branch].append(seller)
    for branch_number, branch in enumerate(branch_names, start=1):
        seller_name = f"seller_options_{branch_number}"
        seller_names = sorted(set(seller_groups[branch]))
        column_letter = get_column_letter(list_column)
        for row, seller in enumerate(seller_names, start=2):
            lists.cell(row=row, column=list_column, value=seller)
        workbook.defined_names.add(
            DefinedName(seller_name, attr_text=f"'Listas tipificacion'!${column_letter}$2:${column_letter}${max(2, len(seller_names) + 1)}")
        )
        lists.cell(row=branch_number + 1, column=9, value=branch)
        lists.cell(row=branch_number + 1, column=10, value=seller_name)
        list_column += 1
    lists.sheet_state = "hidden"

    for column, formula in (
        (activity_column, "=activity_options"),
        (concept_column, f'=INDIRECT(IFERROR(VLOOKUP(${activity_column}2,\'Listas tipificacion\'!$A$2:$B${max(2, len(activity_names) + 1)},2,FALSE),"empty_list"))'),
        (type_column, f'=INDIRECT(IFERROR(VLOOKUP(${activity_column}2&"|"&${concept_column}2,\'Listas tipificacion\'!$D$2:$E${max(2, type_range_number + 1)},2,FALSE),"empty_list"))'),
    ):
        validation = DataValidation(type="list", formula1=formula, allow_blank=allow_blank)
        movements.add_data_validation(validation)
        validation.add(f"{column}2:{column}5000")

    if branch_column and seller_column:
        branch_validation = DataValidation(
            type="list",
            formula1=f"'Sucursales'!$A$2:$A${max(2, len(branch_names) + 1)}",
            allow_blank=allow_blank,
        )
        seller_validation = DataValidation(
            type="list",
            formula1=f'=INDIRECT(IFERROR(VLOOKUP(${branch_column}2,\'Listas tipificacion\'!$I$2:$J${max(2, len(branch_names) + 1)},2,FALSE),"empty_list"))',
            allow_blank=allow_blank,
        )
        movements.add_data_validation(branch_validation)
        movements.add_data_validation(seller_validation)
        branch_validation.add(f"{branch_column}2:{branch_column}5000")
        seller_validation.add(f"{seller_column}2:{seller_column}5000")
