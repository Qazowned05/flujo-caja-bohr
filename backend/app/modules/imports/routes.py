import io
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Annotated, Any, Literal

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import JSONResponse, Response
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.modules.categorias.models import Actividad, Concepto, Tipo
from app.modules.cuentas_bancos.models import CuentaBancaria
from app.modules.imports.models import ImportBatch
from app.modules.imports.schemas import ImportBatchRead, ImportResponse
from app.modules.imports.excel import add_dependent_tipification_lists
from app.modules.shared.audit import log_audit
from app.modules.shared.models import AuditAction
from app.modules.sucursales.models import Sucursal
from app.modules.transacciones.models import (
    EstadoTransaccion,
    OrigenTransaccion,
    Transaccion,
    TransaccionOperacion,
)
from app.modules.usuarios.deps import get_current_user
from app.modules.usuarios.models import User, UserRole
from app.modules.vendedores.models import Vendedor

router = APIRouter(prefix="/imports", tags=["imports"])

HEADERS = [
    "fecha",
    "descripcion",
    "n_operacion",
    "monto",
    "documento",
    "sucursal",
    "vendedor",
    "observaciones",
    "actividad",
    "concepto",
    "tipo",
]
PROJECTION_HEADERS = ["fecha", "descripcion", "monto", "actividad", "concepto", "tipo"]


@router.get("/plantilla.xlsx")
def download_template(
    tipo_importacion: Literal["REAL", "PROYECCION"] = "REAL",
    db: Annotated[Session, Depends(get_db)] = None,
) -> Response:
    workbook = Workbook()
    movements = workbook.active
    movements.title = "Movimientos"
    movements.append(PROJECTION_HEADERS if tipo_importacion == "PROYECCION" else HEADERS)
    if tipo_importacion == "REAL":
        movements.column_dimensions["C"].number_format = "@"
    tipifications = workbook.create_sheet("Tipificaciones")
    tipifications.append(["actividad", "concepto", "tipo"])
    tipification_rows = list(db.execute(
        select(Actividad.nombre, Concepto.nombre, Tipo.nombre)
        .join(Concepto, Tipo.concepto_id == Concepto.id)
        .join(Actividad, Concepto.actividad_id == Actividad.id)
        .where(
            Actividad.is_active.is_(True),
            Concepto.is_active.is_(True),
            Tipo.is_active.is_(True),
        )
        .order_by(Actividad.nombre, Concepto.nombre, Tipo.nombre)
    ))
    for activity, concept, kind in tipification_rows:
        tipifications.append([activity, concept, kind])
    branches = workbook.create_sheet("Sucursales")
    branches.append(["sucursal"])
    branch_names = list(
        db.scalars(select(Sucursal.nombre).where(Sucursal.is_active.is_(True)).order_by(Sucursal.nombre))
    )
    for branch in branch_names:
        branches.append([branch])
    sellers = workbook.create_sheet("Vendedores")
    sellers.append(["vendedor", "sucursal"])
    seller_rows = list(db.execute(
        select(Vendedor.nombre, Sucursal.nombre)
        .outerjoin(Sucursal, Vendedor.sucursal_id == Sucursal.id)
        .where(Vendedor.is_active.is_(True))
        .order_by(Vendedor.nombre)
    ))
    for seller, branch in seller_rows:
        sellers.append([seller, branch or ""])
    activity_column, concept_column, type_column = ("D", "E", "F") if tipo_importacion == "PROYECCION" else ("I", "J", "K")
    add_dependent_tipification_lists(
        workbook,
        movements,
        tipification_rows=tipification_rows,
        branch_names=branch_names,
        seller_rows=seller_rows,
        activity_column=activity_column,
        concept_column=concept_column,
        type_column=type_column,
        branch_column="F" if tipo_importacion == "REAL" else None,
        seller_column="G" if tipo_importacion == "REAL" else None,
        allow_blank=tipo_importacion == "REAL",
    )
    output = io.BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                'attachment; filename="plantilla_proyecciones.xlsx"'
                if tipo_importacion == "PROYECCION"
                else 'attachment; filename="plantilla_importacion.xlsx"'
            )
        },
    )


class CSVValidationError(Exception):
    def __init__(self, errors: list[dict[str, object]]) -> None:
        self.payload = {
            "batch_id": None,
            "total_filas": 0,
            "filas_nuevas": 0,
            "filas_duplicadas": 0,
            "filas_error": len(errors),
            "duplicados": [],
            "errores": errors,
        }


def _parse_date(value: str) -> date | None:
    cleaned = value.strip().strip('"').strip("'").replace("\xa0", "")
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(cleaned, fmt).date()
        except ValueError:
            continue
    return None


def _cell_text(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value or "").strip()


def parse_workbook(content: bytes, headers: list[str], is_projection: bool) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        if "Movimientos" not in workbook.sheetnames:
            raise CSVValidationError([{"fila": 1, "mensaje": "Falta la hoja Movimientos."}])
        sheet = workbook["Movimientos"]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows or [_cell_text(value) for value in rows[0]] != headers:
            raise CSVValidationError([{"fila": 1, "mensaje": "Las cabeceras no son validas."}])
        rows = [dict(zip(headers, row)) for row in rows[1:]]
    except (InvalidFileException, OSError, ValueError) as exc:
        raise CSVValidationError(
            [{"fila": 0, "mensaje": "El archivo debe ser un Excel valido."}]
        ) from exc

    if not rows:
        raise CSVValidationError([{"fila": 2, "mensaje": "El archivo no contiene movimientos."}])

    errors: list[dict[str, object]] = []
    parsed: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows, start=2):
        if None in row or all(not _cell_text(value) for value in row.values()):
            errors.append(
                {"fila": row_number, "mensaje": "No se permiten filas vacias o columnas extra."}
            )
            continue
        values = {key: _cell_text(row.get(key)) for key in headers}
        parsed_date = _parse_date(values["fecha"])
        if not parsed_date:
            errors.append(
                {
                    "fila": row_number,
                    "mensaje": "fecha debe tener formato DD/MM/AAAA (ej: 15/08/2026).",
                }
            )
            continue
        if not values["descripcion"]:
            errors.append(
                {"fila": row_number, "mensaje": "descripcion es obligatoria."}
            )
            continue
        if not is_projection and not values["n_operacion"]:
            errors.append({"fila": row_number, "mensaje": "n_operacion es obligatorio."})
            continue
        if not is_projection and "," in values["n_operacion"] and len(
            operation_numbers(values["n_operacion"])
        ) < 2:
            errors.append(
                {
                    "fila": row_number,
                    "mensaje": "n_operacion multiple requiere al menos dos numeros.",
                }
            )
            continue
        if not is_projection and len(operation_numbers(values["n_operacion"])) != len(
            set(operation_numbers(values["n_operacion"]))
        ):
            errors.append(
                {"fila": row_number, "mensaje": "n_operacion no puede contener numeros repetidos."}
            )
            continue
        if is_projection and not all(
            values[field] for field in ("actividad", "concepto", "tipo")
        ):
            errors.append(
                {
                    "fila": row_number,
                    "mensaje": "actividad, concepto y tipo son obligatorios para proyecciones.",
                }
            )
            continue
        try:
            amount = Decimal(values["monto"].replace(",", "."))
        except (InvalidOperation, ValueError):
            errors.append({"fila": row_number, "mensaje": "monto debe ser un decimal valido."})
            continue
        if not amount.is_finite() or amount == 0:
            errors.append(
                {"fila": row_number, "mensaje": "monto debe ser finito y distinto de cero."}
            )
            continue
        parsed.append(
            {
                "fila": row_number,
                "fecha": parsed_date,
                "descripcion": values["descripcion"],
                "n_operacion": values.get("n_operacion") or None,
                "operaciones": operation_numbers(values.get("n_operacion") or ""),
                "monto": amount,
                "documento": values.get("documento") or None,
                "sucursal": values.get("sucursal") or None,
                "vendedor": values.get("vendedor") or None,
                "observaciones": values.get("observaciones") or None,
                "actividad": values.get("actividad") or None,
                "concepto": values.get("concepto") or None,
                "tipo": values.get("tipo") or None,
            }
        )
    if errors:
        raise CSVValidationError(errors)
    return parsed


def _active_by_name(db: Session, model: Any, name: str) -> Any | None:
    records = list(
        db.scalars(select(model).where(model.nombre == name, model.is_active.is_(True)))
    )
    if len(records) != 1:
        return None
    return records[0]


def is_itf_movement(description: str) -> bool:
    return "itf" in description.casefold()


def operation_numbers(value: str) -> list[str]:
    return [number.strip() for number in value.split(",") if number.strip()]


@router.post("/excel", response_model=ImportResponse)
def import_excel(
    cuenta_bancaria_id: Annotated[uuid.UUID, Form()],
    archivo: Annotated[UploadFile, File()],
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
    tipo_importacion: Annotated[Literal["REAL", "PROYECCION"], Form()] = "REAL",
) -> ImportResponse:
    try:
        if archivo.content_type and archivo.content_type not in {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        }:
            raise CSVValidationError([{"fila": 0, "mensaje": "El archivo debe ser Excel."}])
        is_projection = tipo_importacion == "PROYECCION"
        rows = parse_workbook(
            archivo.file.read(),
            PROJECTION_HEADERS if is_projection else HEADERS,
            is_projection,
        )
    except CSVValidationError as exc:
        return JSONResponse(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, content=exc.payload)

    account = db.get(CuentaBancaria, cuenta_bancaria_id)
    if not account or not account.is_active:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="Cuenta bancaria activa no encontrada.",
        )

    normal_operations = [
        operation
        for row in rows
        if not is_itf_movement(row["descripcion"])
        for operation in row["operaciones"]
    ]
    existing = (
        set(
            db.scalars(
                select(Transaccion.n_operacion).where(
                    Transaccion.cuenta_bancaria_id == cuenta_bancaria_id,
                    Transaccion.is_itf.is_(False),
                    Transaccion.n_operacion.in_(normal_operations),
                )
            )
        )
        if not is_projection
        else set()
    )
    if not is_projection:
        existing.update(
            db.scalars(
                select(TransaccionOperacion.numero)
                .join(Transaccion)
                .where(
                    Transaccion.cuenta_bancaria_id == cuenta_bancaria_id,
                    TransaccionOperacion.numero.in_(normal_operations),
                )
            )
        )
    duplicates: list[dict[str, object]] = []
    new_rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in rows:
        operations = row["operaciones"]
        if is_projection or is_itf_movement(row["descripcion"]):
            new_rows.append(row)
        elif duplicate := next((item for item in operations if item in existing), None):
            duplicates.append(
                {"fila": row["fila"], "n_operacion": duplicate, "motivo": "base_de_datos"}
            )
        elif duplicate := next((item for item in operations if item in seen), None):
            duplicates.append({"fila": row["fila"], "n_operacion": duplicate, "motivo": "archivo"})
        else:
            seen.update(operations)
            new_rows.append(row)

    batch = ImportBatch(
        cuenta_bancaria_id=cuenta_bancaria_id,
        archivo_nombre=archivo.filename or "importacion.csv",
        usuario_id=current_user.id,
        total_filas=len(rows),
        filas_nuevas=len(new_rows),
        filas_duplicadas=len(duplicates),
        filas_error=0,
    )
    try:
        db.add(batch)
        db.flush()
        log_audit(
            db,
            user=current_user,
            table=ImportBatch.__tablename__,
            record_id=batch.id,
            action=AuditAction.CREATE,
            changes={"archivo_nombre": batch.archivo_nombre, "total_filas": batch.total_filas},
        )
        for row in new_rows:
            categories_ok = row.get("actividad") and row.get("concepto") and row.get("tipo")
            next_estado = EstadoTransaccion.PENDIENTE_TIPIFICAR
            actividad_id = None
            concepto_id = None
            tipo_id = None

            if categories_ok:
                actividad = _active_by_name(db, Actividad, row["actividad"])
                if not actividad:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                        detail=f"Actividad activa no encontrada o duplicada: {row['actividad']}",
                    )
                concepto = db.scalar(
                    select(Concepto).where(
                        Concepto.nombre == row["concepto"],
                        Concepto.actividad_id == actividad.id,
                        Concepto.is_active.is_(True),
                    )
                )
                if not concepto:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                        detail=(
                            f"Concepto activo no encontrado para la actividad "
                            f"'{row['actividad']}': {row['concepto']}"
                        ),
                    )
                tipo = db.scalar(
                    select(Tipo).where(
                        Tipo.nombre == row["tipo"],
                        Tipo.concepto_id == concepto.id,
                        Tipo.is_active.is_(True),
                    )
                )
                if not tipo:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                        detail=(
                            f"Tipo activo no encontrado para el concepto "
                            f"'{row['concepto']}': {row['tipo']}"
                        ),
                    )
                actividad_id = actividad.id
                concepto_id = concepto.id
                tipo_id = tipo.id
                next_estado = EstadoTransaccion.TIPIFICADO

            sucursal_id = None
            if not is_projection and row.get("sucursal"):
                sucursal = _active_by_name(db, Sucursal, row["sucursal"])
                if not sucursal:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                        detail=f"Sucursal activa no encontrada o duplicada: {row['sucursal']}",
                    )
                sucursal_id = sucursal.id

            vendedor_id = None
            if not is_projection and row.get("vendedor"):
                vendedor = _active_by_name(db, Vendedor, row["vendedor"])
                if not vendedor:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                        detail=f"Vendedor activo no encontrado o duplicado: {row['vendedor']}",
                    )
                vendedor_id = vendedor.id

            documento = (row["documento"] or "")[:255] or None
            is_multiple = not is_projection and len(row["operaciones"]) > 1
            transaction = Transaccion(
                fecha=row["fecha"],
                descripcion=row["descripcion"],
                n_operacion=None if is_projection or is_multiple else row["n_operacion"],
                is_itf=not is_projection and is_itf_movement(row["descripcion"]),
                monto=row["monto"],
                moneda=account.moneda,
                documento=None if is_projection else documento,
                observaciones=None if is_projection else row["observaciones"],
                sucursal_id=sucursal_id,
                vendedor_id=vendedor_id,
                actividad_id=actividad_id,
                concepto_id=concepto_id,
                tipo_id=tipo_id,
                cuenta_bancaria_id=account.id,
                origen=(
                    OrigenTransaccion.PROYECCION
                    if is_projection
                    else OrigenTransaccion.MULTIPLE
                    if is_multiple
                    else OrigenTransaccion.IMPORTADO
                ),
                estado=EstadoTransaccion.PROYECTADO if is_projection else next_estado,
                materializado=False,
                import_batch_id=batch.id,
                created_by_id=current_user.id,
                updated_by_id=current_user.id,
                operaciones=(
                    [
                        TransaccionOperacion(numero=number, orden=index)
                        for index, number in enumerate(row["operaciones"])
                    ]
                    if is_multiple
                    else []
                ),
            )
            db.add(transaction)
            db.flush()
            log_audit(
                db,
                user=current_user,
                table=Transaccion.__tablename__,
                record_id=transaction.id,
                action=AuditAction.CREATE,
                changes={
                    "n_operacion": transaction.n_operacion,
                    "operaciones": transaction.numeros_operacion,
                    "monto": transaction.monto,
                },
            )
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT, detail="Conflicto de duplicados concurrente."
        ) from exc
    return ImportResponse(
        batch_id=batch.id,
        total_filas=batch.total_filas,
        filas_nuevas=batch.filas_nuevas,
        filas_duplicadas=batch.filas_duplicadas,
        filas_error=0,
        duplicados=duplicates,
        errores=[],
    )


@router.get("/{batch_id}", response_model=ImportBatchRead)
def get_batch(
    batch_id: uuid.UUID,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[Session, Depends(get_db)],
) -> ImportBatch:
    batch = db.get(ImportBatch, batch_id)
    if not batch or (current_user.rol != UserRole.ADMIN and batch.usuario_id != current_user.id):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Lote de importacion no encontrado."
        )
    return batch
