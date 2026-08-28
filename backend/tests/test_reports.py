import io
import uuid
from datetime import date
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.database import Base, get_db
from app.main import app
from app.modules.categorias.models import Actividad, Concepto, Tipo
from app.modules.cuentas_bancos.models import Banco, CuentaBancaria
from app.modules.reports.routes import parse_date
from app.modules.shared.models import AuditAction, AuditLog
from app.modules.sucursales.models import Sucursal
from app.modules.transacciones.models import EstadoTransaccion, OrigenTransaccion, Transaccion
from app.modules.usuarios.deps import get_current_user
from app.modules.usuarios.models import User, UserRole
from app.modules.vendedores.models import Vendedor

HEADERS = [
    "fecha",
    "descripcion",
    "n_operacion",
    "monto",
    "documento",
    "sucursal",
    "vendedor",
    "observaciones",
    "actividad",
    "concepto",
    "tipo",
    "fuente",
]


@pytest.fixture()
def client_and_session():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    session = Session(engine)
    user = User(
        id=uuid.uuid4(),
        email="asesor@example.com",
        hashed_password="unused",
        nombre="Asesor",
        rol=UserRole.ASESOR,
        is_active=True,
        is_superuser=False,
    )
    session.add(user)
    session.commit()

    def override_db():
        yield session

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = lambda: user
    with TestClient(app) as client:
        yield client, session, user
    app.dependency_overrides.clear()
    session.close()
    engine.dispose()


def setup_transaction(session: Session, user: User, **values):
    bank = Banco(nombre="Banco", codigo="BCP")
    session.add(bank)
    session.flush()
    account = CuentaBancaria(banco_id=bank.id, alias="Principal", numero_cuenta="123", moneda="PEN")
    session.add(account)
    session.flush()
    transaction = Transaccion(
        fecha=date(2026, 8, 1),
        descripcion="Original",
        n_operacion="OP-1",
        monto=Decimal("10.00"),
        moneda="PEN",
        cuenta_bancaria_id=account.id,
        origen=OrigenTransaccion.IMPORTADO,
        estado=EstadoTransaccion.PENDIENTE_TIPIFICAR,
        created_by_id=user.id,
        updated_by_id=user.id,
        **values,
    )
    session.add(transaction)
    session.commit()
    return bank, account, transaction


def upload_workbook(client: TestClient, rows: list[list[object]], confirmar: bool = False):
    workbook = Workbook()
    movements = workbook.active
    movements.title = "Movimientos"
    movements.append(HEADERS)
    for row in rows:
        movements.append(row)
    references = workbook.create_sheet("Tipificaciones")
    references.append(["Esta hoja no se importa"])
    content = io.BytesIO()
    workbook.save(content)
    return client.post(
        "/api/v1/reports/actualizacion-masiva",
        data={"confirmar": str(confirmar).lower()},
        files={
            "archivo": (
                "actualizacion.xlsx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def report_row(**values):
    base = {
        "fecha": "2026-08-02",
        "descripcion": "Actualizada",
        "n_operacion": "OP-1",
        "monto": "25.50",
        "documento": "DOC-1",
        "sucursal": "",
        "vendedor": "",
        "observaciones": "Nota",
        "actividad": "",
        "concepto": "",
        "tipo": "",
        "fuente": "BCP | Principal",
    }
    base.update(values)
    return [base[header] for header in HEADERS]


def test_export_has_tipification_sheet_and_only_pending(client_and_session):
    client, session, user = client_and_session
    _, account, pending = setup_transaction(session, user)
    activity = Actividad(nombre="Ventas")
    session.add(activity)
    session.flush()
    concept = Concepto(actividad_id=activity.id, nombre="Cobros")
    session.add(concept)
    session.flush()
    kind = Tipo(concepto_id=concept.id, nombre="Contado")
    session.add(kind)
    branch = Sucursal(nombre="Lima", codigo="LIM")
    session.add(branch)
    session.flush()
    seller = Vendedor(nombre="Ana", codigo="ANA", sucursal_id=branch.id)
    session.add(seller)
    typed = Transaccion(
        fecha=date(2026, 8, 2),
        descripcion="Tipificada",
        n_operacion="OP-2",
        monto=Decimal("20.00"),
        moneda="PEN",
        cuenta_bancaria_id=account.id,
        origen=OrigenTransaccion.IMPORTADO,
        estado=EstadoTransaccion.TIPIFICADO,
        created_by_id=user.id,
        updated_by_id=user.id,
    )
    session.add(typed)
    typed.is_active = False
    session.commit()

    response = client.get("/api/v1/reports/no-tipificados.xlsx")

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["Movimientos", "Tipificaciones", "Sucursales", "Vendedores"]
    rows = list(workbook["Movimientos"].values)
    assert list(rows[0]) == HEADERS
    assert len(rows) == 2
    assert rows[1][0] == "01/08/2026"
    assert rows[1][2] == pending.n_operacion
    assert rows[1][-1] == "BCP | Principal"
    assert list(workbook["Tipificaciones"].values) == [
        ("actividad", "concepto", "tipo"),
        ("Ventas", "Cobros", "Contado"),
    ]
    assert list(workbook["Sucursales"].values) == [("sucursal",), ("Lima",)]
    assert list(workbook["Vendedores"].values) == [("vendedor", "sucursal"), ("Ana", "Lima")]


def test_export_filters_pending_by_creator(client_and_session):
    client, session, user = client_and_session
    _, account, _ = setup_transaction(session, user)
    other_user = User(
        id=uuid.uuid4(),
        email="otro@example.com",
        hashed_password="unused",
        nombre="Otro asesor",
        rol=UserRole.ASESOR,
        is_active=True,
        is_superuser=False,
    )
    other = Transaccion(
        fecha=date(2026, 8, 2),
        descripcion="Otro pendiente",
        n_operacion="OP-OTRA",
        monto=Decimal("15.00"),
        moneda="PEN",
        cuenta_bancaria_id=account.id,
        origen=OrigenTransaccion.IMPORTADO,
        estado=EstadoTransaccion.PENDIENTE_TIPIFICAR,
        created_by_id=other_user.id,
        updated_by_id=other_user.id,
    )
    session.add_all([other_user, other])
    session.commit()

    own = client.get(f"/api/v1/reports/no-tipificados.xlsx?created_by_id={user.id}")
    other_response = client.get(
        f"/api/v1/reports/no-tipificados.xlsx?created_by_id={other_user.id}"
    )

    own_rows = list(load_workbook(io.BytesIO(own.content), data_only=True)["Movimientos"].values)
    assert own.status_code == 200
    assert [row[2] for row in own_rows[1:]] == ["OP-1"]
    assert other_response.status_code == 403

    user.rol = UserRole.ADMIN
    session.commit()
    other_response = client.get(
        f"/api/v1/reports/no-tipificados.xlsx?created_by_id={other_user.id}"
    )
    other_rows = list(
        load_workbook(io.BytesIO(other_response.content), data_only=True)["Movimientos"].values
    )
    assert other_response.status_code == 200
    assert [row[2] for row in other_rows[1:]] == ["OP-OTRA"]


def test_export_audit_includes_full_fields_and_filters_by_account(client_and_session):
    client, session, user = client_and_session
    _, account, transaction = setup_transaction(session, user)
    activity = Actividad(nombre="Operaciones")
    session.add(activity)
    session.flush()
    concept = Concepto(actividad_id=activity.id, nombre="Cobros")
    session.add(concept)
    session.flush()
    kind = Tipo(concepto_id=concept.id, nombre="Contado")
    session.add(kind)
    projection = Transaccion(
        fecha=date(2026, 8, 2),
        descripcion="Cobro proyectado",
        n_operacion=None,
        monto=Decimal("15.00"),
        moneda="PEN",
        cuenta_bancaria_id=account.id,
        actividad_id=activity.id,
        concepto_id=concept.id,
        tipo_id=kind.id,
        origen=OrigenTransaccion.PROYECCION,
        estado=EstadoTransaccion.PROYECTADO,
        created_by_id=user.id,
        updated_by_id=user.id,
    )
    session.add(projection)
    session.commit()

    response = client.get(
        f"/api/v1/reports/auditoria.xlsx?cuenta_bancaria_id={account.id}&tipo_movimiento=PROYECCIONES"
    )

    assert response.status_code == 200
    rows = list(load_workbook(io.BytesIO(response.content), data_only=True)["Auditoria"].values)
    headers = list(rows[0])
    assert {"fuente", "origen", "es_proyeccion", "registrado_por"} <= set(headers)
    projection_row = next(row for row in rows[1:] if row[1] == "Cobro proyectado")
    assert projection_row[headers.index("fuente")] == "BCP | Principal"
    assert projection_row[headers.index("origen")] == "PROYECCION"
    assert projection_row[headers.index("es_proyeccion")] is True
    assert projection_row[headers.index("actividad")] == "Operaciones"
    assert len(rows) == 2
    assert transaction.id


def test_preview_does_not_persist(client_and_session):
    client, session, user = client_and_session
    _, _, transaction = setup_transaction(session, user)

    response = upload_workbook(client, [report_row()])

    assert response.status_code == 200
    assert response.json()["aplicado"] is False
    assert response.json()["filas_validas"] == 1
    session.refresh(transaction)
    assert transaction.descripcion == "Original"
    assert transaction.monto == Decimal("10.00")


def test_preview_accepts_import_date_format(client_and_session):
    client, session, user = client_and_session
    setup_transaction(session, user)

    response = upload_workbook(client, [report_row(fecha="02/08/2026", monto="25,50")])

    assert response.status_code == 200
    assert response.json()["filas_validas"] == 1


def test_preview_accepts_excel_date_cell(client_and_session):
    client, session, user = client_and_session
    setup_transaction(session, user)

    row = report_row()
    row[0] = date(2026, 8, 2)
    response = upload_workbook(client, [row])

    assert response.status_code == 200
    assert response.json()["filas_validas"] == 1


def test_preview_accepts_date_with_excel_text(client_and_session):
    client, session, user = client_and_session
    setup_transaction(session, user)

    response = upload_workbook(client, [report_row(fecha="Fecha: 02/08/2026;")])

    assert response.status_code == 200
    assert response.json()["filas_validas"] == 1


def test_parse_date_accepts_onlyoffice_two_digit_year_format():
    assert parse_date("08-11-26") == date(2026, 8, 11)


def test_confirm_applies_typification_and_audit(client_and_session):
    client, session, user = client_and_session
    _, _, transaction = setup_transaction(session, user)
    activity = Actividad(nombre="Ventas")
    session.add(activity)
    session.flush()
    concept = Concepto(actividad_id=activity.id, nombre="Cobros")
    session.add(concept)
    session.flush()
    kind = Tipo(concepto_id=concept.id, nombre="Contado")
    session.add(kind)
    session.commit()

    response = upload_workbook(
        client,
        [report_row(actividad="Ventas", concepto="Cobros", tipo="Contado")],
        confirmar=True,
    )

    assert response.status_code == 200
    assert response.json()["aplicado"] is True
    session.refresh(transaction)
    assert transaction.estado == EstadoTransaccion.TIPIFICADO
    assert transaction.actividad_id == activity.id
    audit = session.scalar(
        select(AuditLog).where(
            AuditLog.registro_id == transaction.id, AuditLog.accion == AuditAction.UPDATE
        )
    )
    assert audit is not None
    assert "estado" in audit.cambios


@pytest.mark.parametrize(
    "row",
    [
        report_row(fuente="OTRO | Principal"),
        report_row(actividad="Otra", concepto="Cobros", tipo="Contado"),
    ],
)
def test_invalid_source_or_hierarchy_does_not_apply_anything(client_and_session, row):
    client, session, user = client_and_session
    _, _, transaction = setup_transaction(session, user)
    activity = Actividad(nombre="Ventas")
    other_activity = Actividad(nombre="Otra")
    session.add_all([activity, other_activity])
    session.flush()
    concept = Concepto(actividad_id=activity.id, nombre="Cobros")
    session.add(concept)
    session.flush()
    session.add(Tipo(concepto_id=concept.id, nombre="Contado"))
    session.commit()

    response = upload_workbook(client, [row], confirmar=True)

    assert response.status_code == 422
    assert response.json()["aplicado"] is False
    session.refresh(transaction)
    assert transaction.descripcion == "Original"
    assert transaction.estado == EstadoTransaccion.PENDIENTE_TIPIFICAR
    assert session.scalar(select(AuditLog).where(AuditLog.registro_id == transaction.id)) is None
