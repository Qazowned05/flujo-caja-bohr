import io
import uuid
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core.database import Base, get_db
from app.main import app
from app.modules.ganancias_perdidas.routes import TEMPLATE_HEADERS
from app.modules.usuarios.deps import get_current_user
from app.modules.usuarios.models import User, UserRole


@pytest.fixture()
def client_and_session():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    session = Session(engine)
    user = User(
        id=uuid.uuid4(),
        email="admin@example.com",
        hashed_password="unused",
        nombre="Admin",
        rol=UserRole.ADMIN,
        is_active=True,
        is_superuser=True,
    )
    session.add(user)
    session.commit()

    def override_db():
        yield session

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = lambda: user
    previous = settings.profit_and_loss_enabled
    settings.profit_and_loss_enabled = True
    with TestClient(app) as client:
        yield client, user
    settings.profit_and_loss_enabled = previous
    app.dependency_overrides.clear()
    session.close()
    engine.dispose()


def test_imported_entries_build_a_traceable_profit_loss_summary(client_and_session):
    client, _ = client_and_session
    center = client.post(
        "/api/v1/ganancias-perdidas/centros", json={"codigo": "200", "nombre": "Administracion"}
    )
    assert center.status_code == 201
    income = client.post(
        "/api/v1/ganancias-perdidas/rubros",
        json={"codigo": "VENTAS", "nombre": "Ventas", "naturaleza": "INGRESO"},
    )
    expense = client.post(
        "/api/v1/ganancias-perdidas/rubros",
        json={"codigo": "PERSONAL", "nombre": "Personal", "naturaleza": "GASTO_OPERATIVO"},
    )
    assert income.status_code == expense.status_code == 201

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Asientos"
    sheet.append(TEMPLATE_HEADERS)
    sheet.append(
        ["2026-06", "VENTAS", 1000, "200", "Factura", ""]
    )
    sheet.append(
        ["2026-06", "PERSONAL", 300, "200", "Planilla", ""]
    )
    output = io.BytesIO()
    workbook.save(output)
    response = client.post(
        "/api/v1/ganancias-perdidas/importar.xlsx",
        files={
            "archivo": (
                "junio.xlsx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201
    assert response.json()["filas_nuevas"] == 2

    summary = client.get(
        "/api/v1/ganancias-perdidas/resumen?fecha_desde=2026-06-01&fecha_hasta=2026-06-30"
    )
    assert summary.status_code == 200
    body = summary.json()
    assert body["ingresos_netos"] == "1000.00"
    assert body["gastos_operativos"] == "-300.00"
    assert body["utilidad_neta"] == "700.00"
    assert len(client.get("/api/v1/ganancias-perdidas/asientos").json()) == 2


def test_advisor_cannot_change_profit_loss_catalogues(client_and_session):
    client, user = client_and_session
    user.rol = UserRole.ASESOR
    response = client.post(
        "/api/v1/ganancias-perdidas/centros",
        json={"codigo": "200", "nombre": "Administracion"},
    )
    assert response.status_code == 403


def test_distribution_rule_applies_only_without_an_explicit_center(client_and_session):
    client, _ = client_and_session
    first = client.post(
        "/api/v1/ganancias-perdidas/centros", json={"codigo": "300", "nombre": "Santa Natura"}
    ).json()
    second = client.post(
        "/api/v1/ganancias-perdidas/centros", json={"codigo": "310", "nombre": "FQP"}
    ).json()
    rubro = client.post(
        "/api/v1/ganancias-perdidas/rubros",
        json={"codigo": "PERSONAL", "nombre": "Personal", "naturaleza": "GASTO_OPERATIVO"},
    ).json()
    rule = client.post(
        "/api/v1/ganancias-perdidas/reglas-distribucion",
        json={
            "nombre": "Planilla compartida",
            "rubro_id": rubro["id"],
            "lineas": [
                {"centro_resultado_id": first["id"], "porcentaje": "60"},
                {"centro_resultado_id": second["id"], "porcentaje": "40"},
            ],
        },
    )
    assert rule.status_code == 201
    updated_rule = client.patch(
        f"/api/v1/ganancias-perdidas/reglas-distribucion/{rule.json()['id']}",
        json={
            "nombre": "Planilla compartida",
            "rubro_id": rubro["id"],
            "lineas": [
                {"centro_resultado_id": first["id"], "porcentaje": "60"},
                {"centro_resultado_id": second["id"], "porcentaje": "40"},
            ],
        },
    )
    assert updated_rule.status_code == 200
    shared = client.post(
        "/api/v1/ganancias-perdidas/asientos",
        json={
            "fecha": "2026-06-30",
            "cuenta_contable": "621101",
            "descripcion": "Planilla compartida",
            "rubro_id": rubro["id"],
            "debe": "100",
            "haber": "0",
        },
    )
    assert shared.status_code == 201
    assert shared.json()["es_resultado"] is False
    explicit = client.post(
        "/api/v1/ganancias-perdidas/asientos",
        json={
            "fecha": "2026-06-30",
            "cuenta_contable": "621101",
            "descripcion": "Planilla asignada",
            "rubro_id": rubro["id"],
            "centro_resultado_id": first["id"],
            "debe": "50",
            "haber": "0",
        },
    )
    assert explicit.status_code == 201
    summary = client.get("/api/v1/ganancias-perdidas/resumen").json()
    row = next(item for item in summary["rubros"] if item["codigo"] == "PERSONAL")
    assert row["por_centro"] == {"300": "-110.00", "310": "-40.00"}


def test_template_requires_a_rubro_and_guides_the_user(client_and_session):
    client, _ = client_and_session
    client.post(
        "/api/v1/ganancias-perdidas/centros", json={"codigo": "300", "nombre": "Santa Natura"}
    )
    client.post(
        "/api/v1/ganancias-perdidas/rubros",
        json={"codigo": "VENTAS", "nombre": "Ventas", "naturaleza": "INGRESO"},
    )
    template = client.get("/api/v1/ganancias-perdidas/plantilla.xlsx")
    workbook = load_workbook(io.BytesIO(template.content))
    assert {"Asientos", "Lineas de negocio", "Rubros", "Guia de tipificacion"} <= set(workbook.sheetnames)
    assert [cell.value for cell in workbook["Asientos"][1]] == TEMPLATE_HEADERS
    assert {str(validation.sqref) for validation in workbook["Asientos"].data_validations.dataValidation} == {
        "B2:B5000", "D2:D5000"
    }

    invalid = client.post(
        "/api/v1/ganancias-perdidas/asientos",
        json={
            "fecha": "2026-06-30",
            "cuenta_contable": "701101",
            "descripcion": "Venta sin linea",
            "rubro_id": client.get("/api/v1/ganancias-perdidas/rubros").json()[0]["id"],
            "haber": "100",
        },
    )
    assert invalid.status_code == 422
